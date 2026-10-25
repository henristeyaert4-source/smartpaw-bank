"""Leningen, invoer-btw en jaarafsluitingsdossier.

Gebruik:
  python boekhouding/afsluiting.py [jaar]   → boekhouding/rapporten/SmartPaw_Jaarafsluiting_<jaar>.xlsx
Functies leningen() en invoer() worden ook door kwartaal.py gebruikt.
Alles is een VOORBEREIDING: Pro-count controleert. Ramingen staan als raming gemarkeerd.
"""
import glob, json, sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
B = ROOT / "boekhouding"
F = "Arial"
bold, norm, small = Font(name=F, bold=True), Font(name=F), Font(name=F, size=9, italic=True)
h, hf = Font(name=F, bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0F6E5C")
eur = '#,##0.00 €;-#,##0.00 €;"–"'

# Vaste gegevens die niet uit de bank komen (bron: SmartPaw_Bestaande_Financiering.md, Financien_Bron.md)
KBC_KREDIET = {"naam": "KBC investeringskrediet 420-695088", "bedrag": 19947.0, "start": "2026-02-15", "looptijd_mnd": 36,
               "rente": "±4,1% (bevestigen met kredietvoorstel)"}
VASTE_SCHULDEN = [
    {"naam": "Rekening-courant Henri (voorraad uit inbreng jan 2024)", "bedrag": 14000.0, "bron": "oprichting / inbreng", "te_bevestigen": "voorwaarden, rente?"},
]


def _g():
    return json.loads((ROOT / "data/gecategoriseerd.json").read_text())


def leningen(tot=None):
    """Saldo per lening op datum `tot` (incl.), uit de bankverrichtingen."""
    tot = tot or date.today().isoformat()
    g = [r for r in _g() if r["datum"] <= tot]
    som = lambda cat: round(sum(r["bedrag_eur"] for r in g if r["categorie"] == cat), 2)
    kap = -som("Aflossing KBC-krediet")
    rente = -som("Rente KBC-krediet")
    kaydee_in, kaydee_uit = som("Lening Kaydee (in)"), -som("Lening Kaydee (terug)")
    creazy = [r for r in g if r["categorie"] == "Creazy"]
    btw_voor_creazy = sum(-r["bedrag_eur"] for r in g if r["categorie"] == "Btw" and "103039566368" in (r.get("mededeling") or ""))
    creazy_saldo = round(sum(r["bedrag_eur"] for r in creazy) - btw_voor_creazy, 2)
    rijen = [
        {"lening": KBC_KREDIET["naam"], "opgenomen": KBC_KREDIET["bedrag"], "afgelost": round(kap, 2), "saldo": round(KBC_KREDIET["bedrag"] - kap, 2),
         "rente_betaald": round(rente, 2), "uitleg": f"36 mnd vanaf 15/02/2026, rente {KBC_KREDIET['rente']}; kapitaal en rente apart uit de maandelijkse betalingen"},
        {"lening": "Creazy CommV (vennootschap Kaydee)", "opgenomen": round(sum(r['bedrag_eur'] for r in creazy), 2), "afgelost": round(btw_voor_creazy, 2), "saldo": creazy_saldo,
         "rente_betaald": 0.0, "uitleg": "24/06 €10.000 lening + 25/05 €435,72 (reden?) ; 22/01 €1.176,42 = terugstorting btw die SmartPaw op 18/01 voor Creazy betaalde (verrekend)"},
        {"lening": "Lening Kaydee", "opgenomen": kaydee_in, "afgelost": kaydee_uit, "saldo": round(kaydee_in - kaydee_uit, 2), "rente_betaald": 0.0,
         "uitleg": "€9.000 op 05/02 (Revolut), terug op 18/02. Nieuwe lening met intrest? → bevestigen"},
    ]
    for v in VASTE_SCHULDEN:
        rijen.append({"lening": v["naam"], "opgenomen": v["bedrag"], "afgelost": 0.0, "saldo": v["bedrag"], "rente_betaald": 0.0,
                      "uitleg": f"{v['bron']}; {v['te_bevestigen']}"})
    kas = sorted([r for r in g if r["bron"] == "KBC"], key=lambda r: r["datum"])
    return {"tot": tot, "rijen": rijen, "aflossingen": [(r["datum"], -r["bedrag_eur"]) for r in g if r["categorie"] == "Aflossing KBC-krediet" and r["bedrag_eur"]],
            "renteposten": [(r["datum"], -r["bedrag_eur"]) for r in g if r["categorie"] == "Rente KBC-krediet"]}


def invoer(van=None, tot=None):
    """Invoeraangiftes (art. 41) met maatstaf en btw, gefilterd op aangiftedatum."""
    d = json.loads((B / "invoer.json").read_text())
    out = []
    for a in d["aangiftes"]:
        if van and not (van <= a["aangifte_datum"] <= tot):
            continue
        m = a["maatstaf"]
        raming = m is None and a["goederen_eur"] is not None
        if raming:
            m = round(a["goederen_eur"] + (a["vracht_eur"] or 0) + (a["rechten_eur"] or 0), 2)
        out.append(dict(a, maatstaf_gebruikt=m, btw=round(m * 0.21, 2) if m else None, raming=raming or m is None))
    return out


def _tabel(ws, rij, kop, rijen, breedtes, geld=()):
    for j, k in enumerate(kop, 1):
        c = ws.cell(rij, j, k); c.font = h; c.fill = hf; c.alignment = Alignment(wrap_text=True)
    for i, r in enumerate(rijen, rij + 1):
        for j, v in enumerate(r, 1):
            c = ws.cell(i, j, v); c.font = norm
            if j in geld: c.number_format = eur
    for j, w in enumerate(breedtes, 1):
        ws.column_dimensions[chr(64 + j)].width = w
    return rij + 1, rij + len(rijen)


def main(jaar=None):
    jaar = jaar or date.today().year
    tot = f"{jaar}-12-31"
    g = [r for r in _g() if r["datum"].startswith(str(jaar))]
    wb = Workbook()
    s = wb.active; s.title = "Samenvatting"
    s["A1"] = f"SmartPaw VOF — jaarafsluiting {jaar} (voorbereiding voor Pro-count, stand {date.today().isoformat()})"; s["A1"].font = Font(name=F, bold=True, size=13)

    # Voorraad
    snaps = sorted(glob.glob(str(B / "input/voorraad_*.json")))
    vs = json.loads(Path(snaps[-1]).read_text()) if snaps else {"producten": [], "datum": "?"}
    v = wb.create_sheet("Voorraad")
    v["A1"] = f"Voorraad bij Widem aan landed cost — telling {vs['datum']} (op 31/12 opnieuw tellen; container FurVac 750 komt ±26/11 binnen)"; v["A1"].font = bold
    rij = [[p["naam"], p["voorraad"], p["kostprijs"], None, p.get("kostprijs_bron") or ""] for p in vs["producten"]]
    a, z = _tabel(v, 3, ["Product", "Stuks", "Kostprijs/st", "Waarde", "Bron kostprijs"], rij, [30, 9, 13, 14, 60], geld=(3, 4))
    for i in range(a, z + 1):
        v.cell(i, 4, f"=MAX(B{i},0)*C{i}").number_format = eur
    v.cell(z + 1, 1, "Totaal").font = bold; x = v.cell(z + 1, 4, f"=SUM(D{a}:D{z})"); x.number_format = eur; x.font = bold
    v.cell(z + 3, 1, "Nog toe te voegen op 31/12: goederen onderweg (betaald maar niet ontvangen), filters luchtvracht DDP, voorraad bij Henri/kantoor.").font = small
    voorraad_cel = f"Voorraad!D{z+1}"

    # Leningen
    L = leningen(tot)
    l = wb.create_sheet("Leningen")
    l["A1"] = f"Leningen en rekening-courant op {min(tot, date.today().isoformat())} (uit de bank)"; l["A1"].font = bold
    a, z = _tabel(l, 3, ["Lening", "Opgenomen", "Afgelost", "Saldo", "Rente betaald", "Uitleg / te bevestigen"],
                  [[r["lening"], r["opgenomen"], r["afgelost"], r["saldo"], r["rente_betaald"], r["uitleg"]] for r in L["rijen"]], [44, 13, 13, 13, 13, 80], geld=(2, 3, 4, 5))
    l.cell(z + 1, 1, "Totaal schulden aan financiers").font = bold; x = l.cell(z + 1, 4, f"=SUM(D{a}:D{z})"); x.number_format = eur; x.font = bold
    r0 = z + 3
    l.cell(r0, 1, "KBC-krediet: betalingen (kapitaal / rente)").font = bold
    afl = dict(L["aflossingen"]); rnt = dict(L["renteposten"])
    a2, z2 = _tabel(l, r0 + 1, ["Datum", "Kapitaal", "Rente"], [[d_, afl.get(d_, 0.0), rnt.get(d_, 0.0)] for d_ in sorted(set(afl) | set(rnt))], [44, 13, 13], geld=(2, 3))
    leningen_cel = f"Leningen!D{z+1}"

    # Investeringen (kandidaten)
    inv = [r for r in g if r["bedrag_eur"] <= -250 and r["categorie"] in ("Materiaal", "Software & abonnementen", "Marketing & content", "Administratie & lidmaatschappen")]
    iv = wb.create_sheet("Investeringen")
    iv["A1"] = "Uitgaven ≥ €250 die mogelijk een investering zijn (af te schrijven i.p.v. kost) — Pro-count beslist"; iv["A1"].font = bold
    _tabel(iv, 3, ["Datum", "Tegenpartij", "Bedrag incl. btw", "Categorie", "Investering? (ja/nee)"],
           [[r["datum"], (r.get("tegenpartij") or "")[:40], -r["bedrag_eur"], r["categorie"], ""] for r in sorted(inv, key=lambda r: r["datum"])], [12, 40, 16, 32, 20], geld=(3,))

    # Niet-aftrekbaar
    na = [r for r in g if r["categorie"] in ("Representatie & reizen",)]
    n = wb.create_sheet("Niet-aftrekbaar")
    n["A1"] = "Restaurant, representatie en reizen: btw niet aftrekbaar; kosten slechts deels aftrekbaar (restaurant 69%) — bewijs: met wie en waarom"; n["A1"].font = bold
    a, z = _tabel(n, 3, ["Datum", "Tegenpartij", "Bedrag incl. btw", "Met wie / waarom"],
                  [[r["datum"], (r.get("tegenpartij") or "")[:40], -r["bedrag_eur"], ""] for r in sorted(na, key=lambda r: r["datum"])], [12, 40, 16, 50], geld=(3,))
    n.cell(z + 1, 1, "Totaal").font = bold; x = n.cell(z + 1, 3, f"=SUM(C{a}:C{z})"); x.number_format = eur; x.font = bold

    # Invoer
    iw = wb.create_sheet("Invoer-btw")
    iw["A1"] = "Invoeraangiftes met verlegging (art. 41): maatstaf vak 87, btw vak 57 én 59 (netto 0). Ramingen tot de invoeraangifte er is."; iw["A1"].font = bold
    _tabel(iw, 3, ["Aangifte", "Dossier", "Omschrijving", "Maatstaf", "Btw 21%", "Status"],
           [[x["aangifte_datum"], x["dossier"], x["omschrijving"], x["maatstaf_gebruikt"], x["btw"], x["status"]] for x in invoer() if x["aangifte_datum"].startswith(str(jaar))],
           [12, 16, 40, 14, 14, 50], geld=(4, 5))

    # Te regelen
    t = wb.create_sheet("Te regelen")
    t["A1"] = "Open punten voor de jaarafsluiting"; t["A1"].font = bold
    punten = [
        "Voorraadtelling op 31/12 bij Widem (aantallen bevestigen) + goederen onderweg.",
        "Facturen van 2026 die pas in 2027 betaald worden (Widem december, Compass-saldo FurVac 750 ±$5.249, Pro-count jaarafsluiting).",
        "Btw Q3 (€6.873,24) en Q4 nog te betalen → schuld op 31/12.",
        "Rekening-courant Henri €14.000 en lening Creazy €10.000: voorwaarden en rente op papier zetten.",
        "Lening Kaydee: bestaat er nog een lening met intrest (aflossing vanaf eind nov)? Bedrag en contract.",
        "Creazy €435,72 (25/05): waarvoor?",
        "Investeringen (tabblad): wat wordt afgeschreven?",
        "Restaurantkosten: per ticket noteren met wie en waarom.",
        "Kredietkaart 2025 (±€7.870) zonder detail: kaartuittreksels 2025 aan Pro-count.",
        "Verdeling winst/verlies 70/30 of 50/50 vanaf welke datum.",
    ]
    for i, p_ in enumerate(punten, 3):
        t.cell(i, 1, f"☐ {p_}").font = norm
    t.column_dimensions["A"].width = 120

    # Samenvatting
    kas = json.loads((ROOT / "data/status.json").read_text())["saldi"]
    sam = [("Voorraad aan landed cost (telling " + vs["datum"] + ")", f"={voorraad_cel}"),
           ("Schulden aan financiers (KBC, Creazy, Kaydee, rekening-courant Henri)", f"={leningen_cel}"),
           ("Kas KBC (laatste stand)", sum(k["saldo"] for k in kas if k.get("munt") == "EUR" and "KBC" in (k.get("bron") or ""))),
           ("Restaurant/representatie " + str(jaar) + " (deels niet aftrekbaar)", "='Niet-aftrekbaar'!C" + str(n.max_row)),
           ("Mogelijke investeringen te beoordelen (aantal)", max(iv.max_row - 3, 0))]
    for i, (lbl, val) in enumerate(sam, 3):
        s.cell(i, 1, lbl).font = norm; c = s.cell(i, 2, val); c.font = bold
        if i < 7: c.number_format = eur
    s.cell(9, 1, "Tabbladen: Voorraad · Leningen · Investeringen · Niet-aftrekbaar · Invoer-btw · Te regelen").font = small
    s.column_dimensions["A"].width = 70; s.column_dimensions["B"].width = 18
    out = B / "rapporten"; out.mkdir(exist_ok=True)
    pad = out / f"SmartPaw_Jaarafsluiting_{jaar}.xlsx"
    wb.save(pad)
    print(pad)
    return pad


if __name__ == "__main__":
    main(int(sys.argv[1]) if len(sys.argv) > 1 else None)
